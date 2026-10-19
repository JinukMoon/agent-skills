"""Export a matplotlib figure to PPTX: text-free graphics as picture objects plus every text
element (axis labels, legend entries, annotations, panel tags, colour-bar labels) as an editable
Arial text box placed at matplotlib's own position.

Every embedded graphic is written in BOTH formats: a PNG blip plus an `asvg:svgBlip` extension
pointing at a vector SVG of the same region. That is exactly how PowerPoint stores an "SVG
picture" -- PowerPoint draws the SVG (so the plot stays sharp at any zoom / export size) and any
older reader falls back to the PNG. Side-car flat renders of the whole figure are written too
(`<base>_flat.png` / `<base>_flat.svg`).

Each legend becomes one group: its handles (swatches, markers, hatches -- exact matplotlib ink)
as a picture plus its entries as text boxes, so dragging the group moves the whole legend.

Usage (from a figure script namespace that defines `fig`):
    from fig2pptx import export_figure; export_figure(fig, "output/fig2_c.pptx")
"""
import os, re, io
import matplotlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from matplotlib.colors import to_hex
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"   # MS "SVG blip" extension

SYMBOLS = {r"\pm": "±", r"\sim": "~", r"\sigma": "σ", r"\mu": "μ", r"\geq": "≥",
           r"\leq": "≤", r"\times": "×", r"\AA": "Å", r"\chi": "χ", r"\eta": "η",
           r"\lambda": "λ", r"\Delta": "Δ", r"\Sigma": "Σ", r"\theta": "θ", r"\alpha": "α", r"\beta": "β",
           r"\gamma": "γ", r"\epsilon": "ε", r"\phi": "φ", r"\omega": "ω", r"\rightarrow": "→", r"\leftarrow": "←",
           r"\cdot": "·", r"\dagger": "†", r"\%": "%", r"\infty": "∞", r"\Rightarrow": "⇒", r"\,": " ", r"\;": " ", r"\ ": " "}

def _runs_from_math(s):
    """Tokenise mathtext into runs: (text, italic, script) with script in {None,'sub','sup'}."""
    out = []
    i = 0
    def emit(txt, italic=False, script=None):
        if txt: out.append((txt, italic, script))
    def read_group(j):
        # j points at '{' -> return (content, index after '}')
        depth = 0; k = j
        while k < len(s):
            if s[k] == '{': depth += 1
            elif s[k] == '}':
                depth -= 1
                if depth == 0: return s[j+1:k], k+1
            k += 1
        return s[j+1:], len(s)
    while i < len(s):
        c = s[i]
        if c == '\\':
            m = re.match(r"\\([A-Za-z]+|.)", s[i:])
            cmd = m.group(0); i += len(cmd)
            if cmd[1:].isalpha():                      # LaTeX: whitespace after a control word is swallowed ("\Delta G" -> "ΔG")
                while i < len(s) and s[i] == ' ': i += 1
            if cmd in (r"\mathit", r"\mathrm", r"\mathregular", r"\mathdefault", r"\text", r"\textrm", r"\mathbf"):
                if i < len(s) and s[i] == '{':
                    inner, i = read_group(i)
                    for t, it, sc in _runs_from_math(inner):
                        emit(t, it or cmd == r"\mathit", sc)
            elif cmd in SYMBOLS: emit(SYMBOLS[cmd])
            else: emit(cmd.lstrip("\\"))
        elif c in "_^":
            script = 'sub' if c == '_' else 'sup'; i += 1
            if i < len(s) and s[i] == '{':
                inner, i = read_group(i)
                for t, it, sc in _runs_from_math(inner): emit(t, it, sc or script)
            elif i < len(s) and s[i] == '\\':
                m = re.match(r"\\([A-Za-z]+|.)", s[i:]); cmd = m.group(0); i += len(cmd)
                emit(SYMBOLS.get(cmd, cmd.lstrip("\\")), False, script)
            elif i < len(s):
                emit(s[i], False, script); i += 1
        elif c in "{}": i += 1
        else:
            emit(c); i += 1
    return out

def parse_text(s):
    """Split a matplotlib string into runs; $...$ segments go through the math tokeniser."""
    runs = []
    parts = s.split('$')
    for k, part in enumerate(parts):
        if not part: continue
        if k % 2 == 1: runs.extend(_runs_from_math(part))
        else: runs.append((part, False, None))
    return runs

def _collect(fig, bake_ticks=True, bake_legend=False):
    """bake_ticks=True leaves tick labels in the graphics (rendered with the panel);
    only axis labels, titles, legend entries, annotations and panel tags become text boxes.

    bake_legend=True additionally leaves legend entries in the graphics. Use it when a
    label needs typesetting a PowerPoint text run cannot do, such as a crystallographic
    overbar: PowerPoint draws a combining mark (U+0304/U+0305) displaced onto the next
    glyph, whereas matplotlib mathtext ($\\bar{1}$) centres it on its own character."""
    fig.canvas.draw()
    W = fig.bbox.width; H = fig.bbox.height
    items = []
    def grab(t, group=None):
        if t is None: return
        s = t.get_text()
        if not s or not s.strip() or not t.get_visible(): return
        bb = t.get_window_extent(renderer=fig.canvas.get_renderer())
        if bb.width <= 0 or bb.height <= 0: return
        items.append(dict(text=s, cx=(bb.x0 + bb.x1) / 2 / W, cy=1 - (bb.y0 + bb.y1) / 2 / H,
                          w=bb.width / W, h=bb.height / H, size=t.get_fontsize(), rot=t.get_rotation(),
                          weight=t.get_fontweight(), style=t.get_fontstyle(), ha=t.get_ha(),
                          color=to_hex(t.get_color()),    # carry the artist colour (white in-bar labels etc.)
                          group=group))                   # id() of the legend this entry belongs to
        t.set_visible(False)
    def grab_ticks(axis):
        if not axis.get_visible(): return
        lo, hi = sorted(axis.get_view_interval())
        eps = 1e-9 * max(abs(hi - lo), 1e-30)
        for tick in axis.get_major_ticks():
            loc = tick.get_loc()
            if not (lo - eps <= loc <= hi + eps): continue
            for lab in (tick.label1, tick.label2):
                if lab.get_visible(): grab(lab)
    for ax in fig.axes:
        for t in list(ax.texts): grab(t)
        grab(ax.title)
        if ax.axison:
            grab(ax.xaxis.label); grab(ax.yaxis.label)
            if not bake_ticks:
                grab_ticks(ax.xaxis); grab_ticks(ax.yaxis)
    if not bake_legend:
        for leg in _legends(fig):
            for t in leg.get_texts(): grab(t, id(leg))
    for t in list(fig.texts): grab(t)
    return items


def _legends(fig):
    """Visible axes legends and figure legends."""
    legs = [ax.get_legend() for ax in fig.axes if ax.get_visible()] + list(fig.legends)
    return [l for l in legs if l is not None and l.get_visible()]


def _pin_fig_legends(fig):
    """savefig(bbox_inches=<crop>) temporarily shrinks fig.bbox to the crop, and a figure legend
    anchored to fig.bbox re-lays itself out inside every cropped render (a stray swatch in each
    component SVG). Pin each one to its current spot in transFigure coordinates, which the crop
    compensates for. Returns an undo callable."""
    saved = []
    fig.canvas.draw()
    for leg in fig.legends:
        saved.append((leg, leg._loc, leg._bbox_to_anchor, leg.borderaxespad))
        box = leg.get_window_extent().transformed(fig.transFigure.inverted())
        leg.set_bbox_to_anchor(box, transform=fig.transFigure)
        leg._loc = 3                                 # lower left of the pinned box
        leg.borderaxespad = 0
    def undo():
        for leg, loc, anchor, pad in saved:
            # raw attribute: set_bbox_to_anchor() would read the stored pixel box as fractions
            leg._loc = loc; leg._bbox_to_anchor = anchor; leg.borderaxespad = pad; leg.stale = True
    return undo


def save_pptx(prs, path):
    """Save; if the target is locked (open in PowerPoint) fall back to <name>_v2.pptx, _v3, ...
    Returns the path actually written."""
    root, ext = os.path.splitext(path)
    for k in range(1, 100):
        target = path if k == 1 else f"{root}_v{k}{ext}"
        try:
            prs.save(target)
        except PermissionError:
            continue
        if k > 1:
            print(f"    !! {os.path.basename(path)} is locked (open in PowerPoint?) -> saved as {os.path.basename(target)}")
        return target
    raise PermissionError(f"could not write {path} or any _vN variant")


def _check_print_size(items, fig, bake_ticks, Wi, print_width_in, min_pt):
    """Warn about text that lands below min_pt once the figure is printed print_width_in wide."""
    scale = print_width_in / Wi
    sizes = [(it["size"], it["text"]) for it in items]
    if bake_ticks:
        for ax in fig.axes:
            if not (ax.get_visible() and ax.axison): continue
            for lab in ax.get_xticklabels() + ax.get_yticklabels():
                if lab.get_visible() and lab.get_text().strip():
                    sizes.append((lab.get_fontsize(), "tick " + lab.get_text()))
    small = sorted({(round(sz * scale, 1), txt) for sz, txt in sizes if sz * scale < min_pt})
    if small:
        print(f"    !! {len(small)} text(s) below {min_pt} pt at {print_width_in} in print width "
              f"(scale {scale:.2f}):")
        for pt, txt in small[:12]:
            print(f"       {pt:4.1f} pt  {txt[:40]!r}")
    return small

def _save_svg(fig, bbox_in, path, fonttype="path", image_dpi=None):
    """Vector render of one region. bbox_in is a Bbox in figure inches.

    fonttype='path' outlines any text still baked into the graphics (tick numbers). That is
    deliberate: 'none' keeps the glyphs as <text> referencing "Helvetica" by name, and PowerPoint
    on Windows has no Helvetica -- the ticks would silently resample to a substitute face. The
    editable text lives in real PowerPoint text boxes anyway, so nothing is lost here.
    """
    with matplotlib.rc_context({"svg.fonttype": fonttype}):
        fig.savefig(path, format="svg", bbox_inches=bbox_in, transparent=True, dpi=image_dpi or "figure")
    return path


def _attach_svg(slide, pic, svg_path, max_mb=12):
    """Add an asvg:svgBlip extension to `pic` so PowerPoint renders the SVG and keeps the PNG
    as the fallback blip. Skipped (PNG only) if the SVG is implausibly heavy."""
    if not os.path.exists(svg_path): return False
    if os.path.getsize(svg_path) > max_mb * 1024 * 1024: return False
    package = slide.part.package
    taken = {str(part.partname) for part in package.iter_parts()}
    n = 1
    while f"/ppt/media/fig2pptx_svg{n}.svg" in taken: n += 1
    partname = PackURI(f"/ppt/media/fig2pptx_svg{n}.svg")
    with open(svg_path, "rb") as fh: blob = fh.read()
    rId = slide.part.relate_to(Part(partname, "image/svg+xml", package, blob), RT.IMAGE)
    blip = pic._element.blipFill.find(qn("a:blip"))
    blip.append(parse_xml(
        f'<a:extLst xmlns:a="{A_NS}"><a:ext uri="{SVG_EXT_URI}">'
        f'<asvg:svgBlip xmlns:asvg="{SVG_NS}" xmlns:r="{R_NS}" r:embed="{rId}"/>'
        f'</a:ext></a:extLst>'))
    return True


def export_figure(fig, out_path, dpi=600, font="Arial", pad_in=0.03, per_axes=True, bake_ticks=True,
                  bake_legend=False, vector=True, svg_fonttype="path", svg_max_mb=12,
                  print_width_in=None, min_pt=6.0):
    """per_axes=True: every axes (panel, colour bar, render) becomes its own picture object so
    components can be moved independently in PowerPoint; overlapping axes (twin axes) are merged.

    print_width_in: printed width of the figure in the journal (single column ~3.3, double ~7).
    When given, every text box and baked tick label that would print below min_pt is listed.
    Returns (path_written, n_textboxes, n_components, n_vector); the path gets a _v2 suffix when
    the target is open in PowerPoint.

    The figure is restored afterwards (dpi, figure-legend anchors, visibility of every non-text
    artist), also when the export fails; its text stays hidden, as it always has."""
    from matplotlib.text import Text
    dpi0 = fig.dpi
    vis = [(a, a.get_visible()) for a in fig.findobj(lambda a: not isinstance(a, Text))]
    # Measure text at the dpi the ink is rendered at. Layout is not dpi-invariant (mathtext extents
    # are pixel-rounded, so e.g. legend rows with subscripts are ~20 % taller at 100 dpi than at
    # 600): positions read at the screen dpi drift away from the 600-dpi graphics.
    fig.set_dpi(dpi)
    unpin = _pin_fig_legends(fig)
    try:
        return _export(fig, out_path, dpi, dpi0, font, pad_in, per_axes, bake_ticks, bake_legend,
                       vector, svg_fonttype, svg_max_mb, print_width_in, min_pt)
    finally:
        for a, v in vis: a.set_visible(v)
        unpin()
        fig.set_dpi(dpi0)


def _export(fig, out_path, dpi, dpi0, font, pad_in, per_axes, bake_ticks, bake_legend,
            vector, svg_fonttype, svg_max_mb, print_width_in, min_pt):
    from matplotlib.transforms import Bbox
    from PIL import Image
    fig.canvas.draw()
    r = fig.canvas.get_renderer()
    tb = fig.get_tightbbox(r)                       # inches, includes all text
    tb = Bbox.from_extents(tb.x0 - pad_in, tb.y0 - pad_in, tb.x1 + pad_in, tb.y1 + pad_in)
    Wi, Hi = tb.width, tb.height
    base = out_path.replace(".pptx", "")
    fig.savefig(base + "_flat.png", dpi=dpi, bbox_inches=tb)        # complete figure, text baked in
    if vector: _save_svg(fig, tb, base + "_flat.svg", svg_fonttype, dpi0)  # same, vector
    nvec = 0
    items = _collect(fig, bake_ticks=bake_ticks, bake_legend=bake_legend)  # hides text; positions in full-canvas fraction
    fw, fh = fig.get_size_inches()
    for it in items:                                # remap to the tight canvas
        x_in = it["cx"] * fw; y_in = (1 - it["cy"]) * fh
        it["cx"] = (x_in - tb.x0) / Wi; it["cy"] = 1 - (y_in - tb.y0) / Hi
        it["w"] = it["w"] * fw / Wi; it["h"] = it["h"] * fh / Hi
    if print_width_in:
        _check_print_size(items, fig, bake_ticks, Wi, print_width_in, min_pt)
    legends = _legends(fig) if (per_axes and not bake_legend) else []
    leg_pics = {}                                   # id(legend) -> picture shape
    prs = Presentation()
    prs.slide_width = Inches(max(Wi, 1.0)); prs.slide_height = Inches(max(Hi, 1.0))   # PowerPoint minimum 1 in
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if per_axes:
        def _box(a):
            bb = a.get_tightbbox(r)
            if bb is None or bb.width <= 0: return None
            return [bb.x0 / fig.dpi - tb.x0, bb.y0 / fig.dpi - tb.y0, bb.x1 / fig.dpi - tb.x0, bb.y1 / fig.dpi - tb.y0]
        # inset (child) axes become their own components: render them in a second pass
        children = [c for ax in fig.axes for c in ax.child_axes if c.get_visible()]
        child_boxes = [b for b in (_box(c) for c in children) if b is not None]
        # legends become their own components too: box them now, keep them out of the panel renders
        pairs = [(l, b) for l in legends if (b := _box(l)) is not None]
        for l in legends: l.set_visible(False)
        legends = [l for l, _ in pairs]; leg_boxes = [b for _, b in pairs]
        for c in children: c.set_visible(False)
        full = base + "_graphics_full.png"
        fig.savefig(full, dpi=dpi, bbox_inches=tb, transparent=True)   # text-free, transparent, no insets
        im = Image.open(full).convert("RGBA")
        sx, sy = im.size[0] / Wi, im.size[1] / Hi                        # px per inch on the tight canvas
        im2 = None
        if children:
            for c in children: c.set_visible(True)
            full2 = base + "_graphics_insets.png"
            fig.savefig(full2, dpi=dpi, bbox_inches=tb, transparent=True)
            im2 = Image.open(full2).convert("RGBA")
        # axes boxes in tight-canvas inches (tight bbox of each axes now that its text is hidden)
        boxes = []
        for ax in fig.axes:
            if not ax.get_visible(): continue
            for c in ax.child_axes: c.set_visible(False)
            b = _box(ax)
            for c in ax.child_axes: c.set_visible(True)
            if b is not None: boxes.append(b)
        # merge overlapping boxes (twin axes, insets)
        merged = []
        for b in boxes:
            for m in merged:
                if not (b[2] < m[0] or b[0] > m[2] or b[3] < m[1] or b[1] > m[3]):
                    m[0], m[1], m[2], m[3] = min(m[0], b[0]), min(m[1], b[1]), max(m[2], b[2]), max(m[3], b[3]); break
            else:
                merged.append(list(b))
        def _crops(boxes, image, tag, margin):
            out = []
            for k, (x0, y0, x1, y1) in enumerate(boxes):
                X0, X1 = max(0.0, x0 - margin), min(Wi, x1 + margin)
                Y0, Y1 = max(0.0, y0 - margin), min(Hi, y1 + margin)
                crop = image.crop((int(X0 * sx), int((Hi - Y1) * sy), int(X1 * sx), int((Hi - Y0) * sy)))
                cp = f"{base}_{tag}{k:02d}.png"; crop.save(cp)
                out.append((cp, X0, Y0, X1, Y1))
            return out
        # SVG twin of a component must be rendered with the same inset visibility as its PNG:
        # the main components came from `im` (insets hidden), the inset components from `im2`.
        def _svgs(geoms, tag, insets_visible):
            if not vector: return [None] * len(geoms)
            for c in children: c.set_visible(insets_visible)
            paths = []
            for k, (cp, X0, Y0, X1, Y1) in enumerate(geoms):
                sp = f"{base}_{tag}{k:02d}.svg"
                _save_svg(fig, Bbox.from_extents(X0 + tb.x0, Y0 + tb.y0, X1 + tb.x0, Y1 + tb.y0),
                          sp, svg_fonttype, dpi0)
                paths.append(sp)
            for c in children: c.set_visible(True)
            return paths
        main_g = _crops(merged, im, "c", 0.02)
        inset_g = _crops(child_boxes, im2, "i", 0.01) if child_boxes else []
        main_s = _svgs(main_g, "c", False)
        inset_s = _svgs(inset_g, "i", True)
        leg_g, leg_s = [], []
        if legends:
            # third pass: nothing but the legend handles (their text is already hidden)
            hidden = [a for ax in fig.axes for a in ax.get_children() if a.get_visible()]
            hidden += [a for a in fig.texts + fig.lines + fig.patches + fig.images + fig.artists
                       if a.get_visible()]
            for a in hidden: a.set_visible(False)
            for l in legends: l.set_visible(True)
            full3 = base + "_graphics_legends.png"
            fig.savefig(full3, dpi=dpi, bbox_inches=tb, transparent=True)
            leg_g = _crops(leg_boxes, Image.open(full3).convert("RGBA"), "l", 0.03)
            if vector:
                for k, (cp, X0, Y0, X1, Y1) in enumerate(leg_g):
                    sp = f"{base}_l{k:02d}.svg"
                    _save_svg(fig, Bbox.from_extents(X0 + tb.x0, Y0 + tb.y0, X1 + tb.x0, Y1 + tb.y0),
                              sp, svg_fonttype, dpi0)
                    leg_s.append(sp)
            else:
                leg_s = [None] * len(leg_g)
            for a in hidden: a.set_visible(True)
            for l in legends: l.set_visible(False)
        for (cp, X0, Y0, X1, Y1), sp in list(zip(main_g, main_s)) + list(zip(inset_g, inset_s)):
            pic = slide.shapes.add_picture(cp, Inches(X0), Inches(Hi - Y1),
                                           width=Inches(X1 - X0), height=Inches(Y1 - Y0))
            if sp: nvec += _attach_svg(slide, pic, sp, svg_max_mb)
        for l, (cp, X0, Y0, X1, Y1), sp in zip(legends, leg_g, leg_s):
            pic = slide.shapes.add_picture(cp, Inches(X0), Inches(Hi - Y1),
                                           width=Inches(X1 - X0), height=Inches(Y1 - Y0))
            if sp: nvec += _attach_svg(slide, pic, sp, svg_max_mb)
            leg_pics[id(l)] = pic
        ncomp = len(merged) + len(child_boxes) + len(leg_g)
    else:
        png = base + "_graphics.png"
        fig.savefig(png, dpi=dpi, bbox_inches=tb, transparent=True)
        pic = slide.shapes.add_picture(png, 0, 0, width=Inches(Wi), height=Inches(Hi)); ncomp = 1
        if vector:
            svg = _save_svg(fig, tb, base + "_graphics.svg", svg_fonttype, dpi0)
            nvec += _attach_svg(slide, pic, svg, svg_max_mb)
    leg_boxes_tb = {}                               # id(legend) -> its text boxes
    for it in items:
        rot = it["rot"] % 360
        w_in, h_in = it["w"] * Wi, it["h"] * Hi
        if abs(rot - 90) < 1 or abs(rot - 270) < 1:
            w_in, h_in = h_in, w_in
        pad = 0.06
        w_in += pad; h_in += pad * 0.5
        cx, cy = it["cx"] * Wi, it["cy"] * Hi
        box = slide.shapes.add_textbox(Inches(cx - w_in / 2), Inches(cy - h_in / 2), Inches(w_in), Inches(h_in))
        box.rotation = (-rot) % 360
        tf = box.text_frame; tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = it["text"].split("\n")
        for li, line in enumerate(lines):
            p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}.get(it["ha"], PP_ALIGN.CENTER)
            for txt, italic, script in parse_text(line):
                rr = p.add_run(); rr.text = txt
                f = rr.font; f.name = font; f.size = Pt(it["size"]); f.italic = italic or (it["style"] == "italic")
                f.bold = it["weight"] in ("bold", "semibold", "heavy")
                f.color.rgb = RGBColor.from_string(it.get("color", "#000000").lstrip("#")[:6].upper())
                if script == "sub": rr.font._element.set("baseline", "-25000")
                elif script == "sup": rr.font._element.set("baseline", "30000")
        if it.get("group") in leg_pics:
            leg_boxes_tb.setdefault(it["group"], []).append(box)
    for gid, pic in leg_pics.items():               # legend = handles picture + entry text boxes, one group
        grp = slide.shapes.add_group_shape([pic] + leg_boxes_tb.get(gid, []))
        grp.name = "legend"
    out_path = save_pptx(prs, out_path)
    return out_path, len(items), ncomp, nvec
