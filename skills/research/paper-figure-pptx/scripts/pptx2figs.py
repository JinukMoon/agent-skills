"""PPTX (the editable master) -> PNG + SVG. Run this after hand-editing a figure deck.

The PPTX is the source of truth; PNG and SVG are regenerated from whatever it currently
contains, so editing text in PowerPoint and re-running is the whole workflow.

  PNG : rendered by real PowerPoint (COM), so it is pixel-identical to what you see on screen.
  SVG : composed here. PowerPoint cannot export SVG (no graphics filter is
        installed) and its PDF export resamples pictures to ~101 ppi, so neither is usable.
        Instead each picture's embedded vector part (written by fig2pptx.py as an
        `asvg:svgBlip`) is inlined verbatim and every text box becomes a real <text> element.
        The result is fully vector and its text stays selectable/editable in
        Illustrator/Inkscape.

Usage:
    python pptx2figs.py deck.pptx [more.pptx ...] [--png-width 4800] [--no-png] [--no-svg]
                        [--print-width 3.3 [--min-pt 6]]
"""
import argparse, base64, os, re, shutil, subprocess, sys, tempfile

from pptx import Presentation
from pptx.oxml.ns import qn

EMU_PT = 12700.0                      # 1 pt = 12700 EMU
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
WIN_PS = "/mnt/c/Windows/System32/WindowsPowerShell/v1.0"
# PowerPoint is single-instance: New-Object attaches to the user's running PowerPoint if there is
# one, so an unconditional Quit() would close the user's open decks. Quit only when no PowerPoint
# window is on screen (i.e. the instance is the invisible one this script started).
PP_CLOSE = ('$pres.Close();'
            'if (-not (Get-Process POWERPNT -ErrorAction SilentlyContinue | '
            'Where-Object { $_.MainWindowTitle })) { $pp.Quit() }')


def _win_workdir():
    """Fresh work dir under the Windows %TEMP% -> (wsl_path, win_path). PowerPoint must see the file."""
    os.environ["PATH"] = os.environ.get("PATH", "") + ":" + WIN_PS
    r = subprocess.run(["powershell.exe", "-NoProfile", "-Command", "$env:TEMP"],
                       capture_output=True, text=True)
    win_temp = r.stdout.strip()
    if r.returncode != 0 or not win_temp:
        raise RuntimeError("could not read Windows %TEMP% via powershell.exe (WSL + PowerPoint required)")
    wsl_temp = subprocess.run(["wslpath", "-u", win_temp], capture_output=True, text=True, check=True).stdout.strip()
    work = tempfile.mkdtemp(prefix="pptx2figs_", dir=wsl_temp)
    win = subprocess.run(["wslpath", "-w", work], capture_output=True, text=True, check=True).stdout.strip()
    return work, win


# ----------------------------------------------------------------- PNG via PowerPoint
def render_png(pptx_path, out_png, width_px=4800):
    """Drive the real PowerPoint to export slide 1. Needs the file on a Windows-visible path."""
    work, win = _win_workdir()
    try:
        shutil.copy(pptx_path, os.path.join(work, "in.pptx"))
        ps = (
            '$pp=New-Object -ComObject PowerPoint.Application;'
            f'$pres=$pp.Presentations.Open("{win}\\in.pptx",$true,$false,$false);'
            f'$pres.Slides.Item(1).Export("{win}\\out.png","PNG",{int(width_px)},0);'
            + PP_CLOSE
        )
        subprocess.run(["powershell.exe", "-NoProfile", "-Command", ps],
                       check=True, capture_output=True)
        src = os.path.join(work, "out.png")
        if not os.path.exists(src):
            raise RuntimeError("PowerPoint produced no PNG")
        shutil.copy(src, out_png)
        return out_png
    finally:
        shutil.rmtree(work, ignore_errors=True)


def svg_via_powerpoint_pdf(pptx_path, out_svg):
    """Fallback for decks holding shapes the composer cannot draw (lines, autoshapes, groups):
    let PowerPoint write a PDF, then vectorise it with pdftocairo.

    Lines, autoshapes and text survive as vector. Pictures do NOT: PowerPoint resamples them to
    ~200 ppi JPEG on the way into the PDF (and to only ~100 ppi if they are SVG pictures), so the
    PowerPoint-rendered PNG stays the better deliverable whenever a figure leans on its images.
    """
    work, win = _win_workdir()
    try:
        shutil.copy(pptx_path, os.path.join(work, "in.pptx"))
        ps = ('$pp=New-Object -ComObject PowerPoint.Application;'
              f'$pres=$pp.Presentations.Open("{win}\\in.pptx",$true,$false,$false);'
              f'$pres.SaveCopyAs("{win}\\out.pdf", 32);'
              + PP_CLOSE)
        subprocess.run(["powershell.exe", "-NoProfile", "-Command", ps],
                       check=True, capture_output=True)
        pdf = os.path.join(work, "out.pdf")
        if not os.path.exists(pdf):
            raise RuntimeError("PowerPoint produced no PDF")
        subprocess.run(["pdftocairo", "-svg", "-f", "1", "-l", "1", pdf, out_svg], check=True)
        return out_svg
    finally:
        shutil.rmtree(work, ignore_errors=True)


# ----------------------------------------------------------------- SVG composer
def _ns_ids(svg_text, prefix):
    """Make every id in one component unique. Ids are document-global in SVG, so two
    matplotlib components inlined side by side would otherwise clobber each other's
    clip paths and marker definitions."""
    for i in sorted(set(re.findall(r'id="([^"]+)"', svg_text)), key=len, reverse=True):
        svg_text = (svg_text.replace(f'id="{i}"', f'id="{prefix}{i}"')
                            .replace(f'#{i}"', f'#{prefix}{i}"')
                            .replace(f'url(#{i})', f'url(#{prefix}{i})'))
    return svg_text


def _inline_picture(blob, prefix, x, y, w, h):
    t = blob.decode("utf-8", "replace")
    t = re.sub(r"<\?xml[^>]*\?>", "", t)
    t = re.sub(r"<!DOCTYPE.*?>", "", t, flags=re.S)
    t = re.sub(r"<metadata>.*?</metadata>", "", t, flags=re.S)
    t = _ns_ids(t, prefix)
    m = re.search(r"<svg\b[^>]*>", t)
    if not m:
        return None
    head = m.group(0)
    vb = re.search(r'viewBox="([^"]+)"', head)
    keep = f' viewBox="{vb.group(1)}"' if vb else ""
    new = (f'<svg x="{x:.4f}" y="{y:.4f}" width="{w:.4f}" height="{h:.4f}"{keep} '
           f'overflow="hidden" preserveAspectRatio="none">')
    return new + t[m.end():]


def _raster_picture(blob, ext, x, y, w, h):
    href = f"data:image/{'png' if ext == 'png' else ext};base64," + base64.b64encode(blob).decode()
    return (f'<image x="{x:.4f}" y="{y:.4f}" width="{w:.4f}" height="{h:.4f}" '
            f'preserveAspectRatio="none" xlink:href="{href}"/>')


def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _text_svg(shape, x0, y0, w, h):
    """A text box -> one <text> per paragraph, positioned the way PowerPoint lays it out:
    horizontally per the paragraph alignment, vertically centred in the box (fig2pptx sets
    MIDDLE anchor and switches word wrap off). (x0, y0, w, h) is the box on the slide, in pt."""
    cx, cy = x0 + w / 2, y0 + h / 2
    paras = [p for p in shape.text_frame.paragraphs]
    sizes = []
    for p in paras:
        s = [r.font.size.pt for r in p.runs if r.font.size is not None]
        sizes.append(max(s) if s else 12.0)
    lh = [s * 1.2 for s in sizes]
    total = sum(lh)
    out = []
    y = cy - total / 2
    for p, size, step in zip(paras, sizes, lh):
        if not p.runs:
            y += step
            continue
        align = p.alignment.name if p.alignment is not None else "CENTER"   # str() gives "LEFT (1)"
        anchor, ax = {"LEFT": ("start", x0), "RIGHT": ("end", x0 + w)}.get(align, ("middle", cx))
        base = y + step / 2 + size * 0.35          # baseline of a vertically centred line
        spans = []
        for r in p.runs:
            f = r.font
            st = [f'font-family:{f.name or "Arial"}', f'font-size:{(f.size.pt if f.size else size):.2f}px']
            if f.bold: st.append("font-weight:bold")
            if f.italic: st.append("font-style:italic")
            try:
                if f.color and f.color.rgb is not None: st.append(f"fill:#{f.color.rgb}")
            except (AttributeError, TypeError):
                pass
            bl = r.font._element.get("baseline")
            extra = ""
            if bl:
                shift = int(bl) / 1000.0                      # 1000ths of a percent
                extra = f' baseline-shift="{shift:.0f}%" font-size="65%"'
            spans.append(f'<tspan style="{";".join(st)}"{extra}>{_esc(r.text)}</tspan>')
        out.append(f'<text x="{ax:.4f}" y="{base:.4f}" text-anchor="{anchor}" '
                   f'xml:space="preserve">{"".join(spans)}</text>')
        y += step
    if shape.rotation:
        # PowerPoint rotation is clockwise degrees about the shape centre; so is SVG's rotate()
        return [f'<g transform="rotate({shape.rotation:.4f} {cx:.4f} {cy:.4f})">'] + out + ["</g>"]
    return out


def _group_map(group, parent):
    """Child-space -> slide-space map of a group, composed with the parent map.
    A map is (ax, bx, ay, by) with slide = a * child + b, all in EMU."""
    xfrm = group._element.grpSpPr.find(qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    sx = int(ext.get("cx")) / max(int(ch_ext.get("cx")), 1)
    sy = int(ext.get("cy")) / max(int(ch_ext.get("cy")), 1)
    bx = int(off.get("x")) - sx * int(ch_off.get("x"))
    by = int(off.get("y")) - sy * int(ch_off.get("y"))
    pax, pbx, pay, pby = parent
    return pax * sx, pax * bx + pbx, pay * sy, pay * by + pby


def _rect(shape, m):
    """Shape rectangle on the slide in pt, through map m."""
    ax, bx, ay, by = m
    return ((ax * shape.left + bx) / EMU_PT, (ay * shape.top + by) / EMU_PT,
            ax * shape.width / EMU_PT, ay * shape.height / EMU_PT)


def iter_shapes(shapes, m=(1.0, 0.0, 1.0, 0.0)):
    """Yield (shape, map) for every leaf shape, descending into groups. A rotated or flipped group
    is yielded whole (map None): the composer cannot draw it and routes the SVG through the PDF."""
    for shape in shapes:
        if shape._element.tag == qn("p:grpSp"):
            xfrm = shape._element.grpSpPr.find(qn("a:xfrm"))
            if int(xfrm.get("rot", 0)) or xfrm.get("flipH") in ("1", "true") or xfrm.get("flipV") in ("1", "true"):
                yield shape, None
                continue
            yield from iter_shapes(shape.shapes, _group_map(shape, m))
        else:
            yield shape, m


def compose_svg(pptx_path, out_svg, slide_index=0):
    prs = Presentation(pptx_path)
    if len(prs.slides) > 1:
        print(f"    note: {len(prs.slides)} slides; only slide {slide_index + 1} is exported")
    W, H = prs.slide_width / EMU_PT, prs.slide_height / EMU_PT
    slide = prs.slides[slide_index]
    body, skipped, n_vec, n_ras = [], [], 0, 0
    for k, (shape, m) in enumerate(iter_shapes(slide.shapes)):
        el = shape._element
        if m is None:
            skipped.append(f"rotated/flipped group #{k}")
        elif el.tag == qn("p:pic"):
            x, y, w, h = _rect(shape, m)
            blip = el.blipFill.find(qn("a:blip"))
            svg_blob = None
            ext = blip.find(qn("a:extLst")) if blip is not None else None
            if ext is not None:
                node = ext.find(f".//{{{SVG_NS}}}svgBlip")
                if node is not None:
                    rid = node.get(qn("r:embed"))
                    if rid:
                        svg_blob = shape.part.related_part(rid).blob
            if svg_blob is not None:
                frag = _inline_picture(svg_blob, f"c{k}_", x, y, w, h)
                if frag:
                    body.append(frag); n_vec += 1; continue
            rid = blip.get(qn("r:embed")) if blip is not None else None
            if rid:
                part = shape.part.related_part(rid)
                body.append(_raster_picture(part.blob, str(part.partname).rsplit(".", 1)[-1].lower(),
                                            x, y, w, h))
                n_ras += 1
            else:
                skipped.append(f"picture #{k} with no image")
        elif shape.has_text_frame and shape.text_frame.text.strip():
            body.extend(_text_svg(shape, *_rect(shape, m)))
        else:
            skipped.append(f"{shape.shape_type} #{k}")
    doc = (f'<?xml version="1.0" encoding="utf-8"?>\n'
           f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
           f'width="{W:.4f}pt" height="{H:.4f}pt" viewBox="0 0 {W:.4f} {H:.4f}" version="1.1">\n'
           + "\n".join(body) + "\n</svg>\n")
    with open(out_svg, "w", encoding="utf-8") as fh:
        fh.write(doc)
    return out_svg, n_vec, n_ras, skipped


def check_print_size(pptx_path, print_width_in, min_pt=6.0, slide_index=0):
    """List text boxes that print below min_pt when the figure is print_width_in wide.
    Baked tick numbers live inside pictures and are checked by fig2pptx at export time."""
    prs = Presentation(pptx_path)
    scale = print_width_in * 72.0 / (prs.slide_width / EMU_PT)
    small = []
    for shape, _ in iter_shapes(prs.slides[slide_index].shapes):
        if not (shape.has_text_frame and shape.text_frame.text.strip()): continue
        sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs
                 if r.font.size is not None and r.font._element.get("baseline") is None]
        if sizes and min(sizes) * scale < min_pt:
            small.append((min(sizes) * scale, shape.text_frame.text.replace("\n", " ")))
    if small:
        print(f"    !! {len(small)} text box(es) below {min_pt} pt at {print_width_in} in print width "
              f"(scale {scale:.2f}):")
        for pt, txt in sorted(small)[:12]:
            print(f"       {pt:4.1f} pt  {txt[:40]!r}")
    return small


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("pptx", nargs="+")
    ap.add_argument("--png-width", type=int, default=4800)
    ap.add_argument("--suffix", default="_final",
                    help="appended to the deck name for the outputs (default _final, so a deck "
                         "built from fig2_c.py does not clobber that script's own fig2_c.png)")
    ap.add_argument("--no-png", action="store_true")
    ap.add_argument("--no-svg", action="store_true")
    ap.add_argument("--print-width", type=float, default=None,
                    help="printed figure width in inches (single column ~3.3, double ~7); "
                         "warns about text boxes that would print below --min-pt")
    ap.add_argument("--min-pt", type=float, default=6.0)
    ap.add_argument("--force-compose", action="store_true",
                    help="keep the composed SVG even if some shapes had to be dropped")
    a = ap.parse_args(argv)
    for p in a.pptx:
        base = os.path.splitext(os.path.abspath(p))[0] + a.suffix
        if a.print_width:
            check_print_size(p, a.print_width, a.min_pt)
        if not a.no_svg:
            out, nv, nr, skipped = compose_svg(p, base + ".svg")
            size = os.path.getsize(out) / 1024
            if skipped and not a.force_compose:
                kinds = sorted({s.split(" #")[0] for s in skipped})
                print(f"{os.path.basename(p):32s} {len(skipped)} shape(s) the composer cannot draw "
                      f"({', '.join(kinds)}) -> routing the SVG through PowerPoint's PDF")
                out = svg_via_powerpoint_pdf(p, base + ".svg")
                print(f"{os.path.basename(out):32s} {os.path.getsize(out)/1024:.0f} KB via PDF; "
                      f"shapes and text are vector, pictures were resampled to ~200 ppi by "
                      f"PowerPoint -- prefer the PNG if this figure leans on its images")
            else:
                print(f"{os.path.basename(out):32s} {nv} vector + {nr} raster pictures, {size:.0f} KB")
                for s in skipped:
                    print(f"    !! left out of the SVG: {s}")
        if not a.no_png:
            out = render_png(p, base + ".png", a.png_width)
            print(f"{os.path.basename(out):32s} {os.path.getsize(out)/1024:.0f} KB "
                  f"({a.png_width} px wide, rendered by PowerPoint)")


if __name__ == "__main__":
    main()
